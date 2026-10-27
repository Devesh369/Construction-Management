"""
Reports generation services.
Generates in-memory PDF (ReportLab) and Excel (openpyxl) documents.
Layer 7: Read-only consumer of workforce, attendance, progress, expenses, payroll, ledger.
"""

import os
from decimal import Decimal
from io import BytesIO

import openpyxl
from django.conf import settings
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from attendance.models import Attendance, AttendanceStatus
from core.exceptions import RowLimitExceededError
from core.money import q2
from expenses.models import Expense
from ledger.services import profit_and_loss, trial_balance
from payroll.models import PayrollVoucher
from progress.models import DailyProgressReport
from projects.models import Project

ROW_LIMIT_CAP = 10000

LETTERHEAD_PORTRAIT = os.path.join(
    os.path.dirname(__file__), "assets", "buildsphere_letterhead.png"
)
LETTERHEAD_LANDSCAPE = os.path.join(
    os.path.dirname(__file__), "assets", "buildsphere_letterhead_landscape.png"
)


def draw_letterhead_background(canvas_obj, doc):
    """
    Draws the official BuildSphere letterhead on every page background.
    Uses landscape-proportioned letterhead for landscape pages and
    portrait letterhead for portrait pages.
    """
    canvas_obj.saveState()
    page_w, page_h = doc.pagesize
    img_path = LETTERHEAD_LANDSCAPE if page_w > page_h else LETTERHEAD_PORTRAIT
    if os.path.exists(img_path):
        canvas_obj.drawImage(
            img_path,
            0,
            0,
            width=page_w,
            height=page_h,
            preserveAspectRatio=False,
        )
    canvas_obj.restoreState()


def get_user_display_name(user) -> str:
    """Safely return user full name or email."""
    if not user:
        return "System"
    return getattr(user, "full_name", "") or getattr(user, "email", str(user))


# ==============================================================================
# ReportLab Canvas with Page Numbers
# ==============================================================================


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass ReportLab canvas that calculates total pages dynamically
    and prints 'Page X of Y' in the running footer.
    """

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, total_pages: int):
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.white)

        page_w = self._pagesize[0]

        # Running footer text cleanly positioned on the official letterhead base wave
        footer_text = f"Page {self._pageNumber} of {total_pages}"
        self.drawRightString(page_w - 36, 16, footer_text)
        site_name = getattr(settings, "SITE_NAME", "BuildSphere")
        site_tagline = getattr(settings, "SITE_TAGLINE", "Construction Management Platform")
        self.drawString(36, 16, f"{site_name} • Confidential • {site_tagline}")

        self.restoreState()


# ==============================================================================
# Excel Styling Helpers
# ==============================================================================


def apply_excel_styling(ws, header_row_index: int = 4):
    """
    Apply clean professional styling to an openpyxl worksheet:
    - Primary navy fill for table header.
    - Thin borders.
    - Auto-adjust column widths.
    """
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Style header row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row_index, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Style data rows
    for row in range(header_row_index + 1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.font = Font(name="Arial", size=9)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


# ==============================================================================
# 1. Payroll Register PDF Export
# ==============================================================================


def generate_payroll_pdf(
    org,
    user,
    project_id=None,
    period_start=None,
    period_end=None,
    allowed_pids=None,
) -> tuple[BytesIO, int]:
    """
    Generate itemized payroll wage sheet PDF.
    """
    qs = PayrollVoucher.objects.filter(
        run__organization=org, deleted_at__isnull=True
    ).select_related("employee", "run", "run__project")

    if project_id:
        qs = qs.filter(run__project_id=project_id)
    elif allowed_pids is not None:
        qs = qs.filter(run__project_id__in=allowed_pids)

    if period_start:
        qs = qs.filter(run__period_start__gte=period_start)
    if period_end:
        qs = qs.filter(run__period_end__lte=period_end)

    row_count = qs.count()
    if row_count > ROW_LIMIT_CAP:
        raise RowLimitExceededError()

    vouchers = list(qs.order_by("run__period_start", "voucher_number"))

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=120,
        bottomMargin=65,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569"),
    )
    cell_style = ParagraphStyle(
        "CellNormal",
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold = ParagraphStyle(
        "CellBold",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#0F172A"),
    )

    story = [
        Paragraph(f"{org.name} - Payroll Register", title_style),
        Spacer(1, 4),
        Paragraph(
            f"Period: {period_start or 'All'} to {period_end or 'All'} | "
            f"Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
            meta_style,
        ),
        Spacer(1, 14),
    ]

    table_data = [
        [
            Paragraph("Voucher #", cell_bold),
            Paragraph("Worker Code", cell_bold),
            Paragraph("Worker Name", cell_bold),
            Paragraph("Project", cell_bold),
            Paragraph("Present / Half", cell_bold),
            Paragraph("OT Hrs", cell_bold),
            Paragraph("Basic Pay", cell_bold),
            Paragraph("OT Pay", cell_bold),
            Paragraph("Gross Pay", cell_bold),
            Paragraph("Advances", cell_bold),
            Paragraph("Net Pay", cell_bold),
        ]
    ]

    tot_basic = Decimal("0.00")
    tot_ot = Decimal("0.00")
    tot_gross = Decimal("0.00")
    tot_adv = Decimal("0.00")
    tot_net = Decimal("0.00")
    tot_pres = 0
    tot_half = 0

    for v in vouchers:
        tot_basic += v.basic_pay
        tot_ot += v.ot_pay
        tot_gross += v.gross_pay
        tot_adv += v.advances_deducted
        tot_net += v.net_pay
        tot_pres += v.days_present
        tot_half += v.half_days

        table_data.append(
            [
                Paragraph(str(v.voucher_number), cell_style),
                Paragraph(str(v.employee.employee_code), cell_style),
                Paragraph(str(v.employee.full_name), cell_style),
                Paragraph(str(v.run.project.code), cell_style),
                Paragraph(f"{v.days_present} / {v.half_days}", cell_style),
                Paragraph(str(v.ot_hours), cell_style),
                Paragraph(str(v.basic_pay), cell_style),
                Paragraph(str(v.ot_pay), cell_style),
                Paragraph(str(v.gross_pay), cell_style),
                Paragraph(str(v.advances_deducted), cell_style),
                Paragraph(str(v.net_pay), cell_style),
            ]
        )

    # Totals Row
    table_data.append(
        [
            Paragraph("TOTALS", cell_bold),
            Paragraph("", cell_bold),
            Paragraph("", cell_bold),
            Paragraph("", cell_bold),
            Paragraph(f"{tot_pres} / {tot_half}", cell_bold),
            Paragraph("", cell_bold),
            Paragraph(str(q2(tot_basic)), cell_bold),
            Paragraph(str(q2(tot_ot)), cell_bold),
            Paragraph(str(q2(tot_gross)), cell_bold),
            Paragraph(str(q2(tot_adv)), cell_bold),
            Paragraph(str(q2(tot_net)), cell_bold),
        ]
    )

    col_widths = [70, 60, 110, 60, 65, 45, 65, 55, 65, 55, 65]
    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0F172A")),
                ("LINEBELOW", (0, -1), (-1, -1), 1.5, colors.HexColor("#0F172A")),
                ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#0F172A")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F8FAFC")),
            ]
        )
    )

    story.append(table)
    doc.build(
        story,
        canvasmaker=NumberedCanvas,
        onFirstPage=draw_letterhead_background,
        onLaterPages=draw_letterhead_background,
    )
    buffer.seek(0)
    return buffer, row_count


# ==============================================================================
# 2. Daily Progress Report (DPR) Excel Export
# ==============================================================================


def generate_dpr_excel(
    org, user, project_id=None, date_from=None, date_to=None, allowed_pids=None
) -> tuple[BytesIO, int]:
    """
    Generate DPR summary with tasks and materials usage Excel log.
    """
    qs = DailyProgressReport.objects.filter(
        organization=org, deleted_at__isnull=True
    ).select_related("project", "submitted_by", "verified_by")

    if project_id:
        qs = qs.filter(project_id=project_id)
    elif allowed_pids is not None:
        qs = qs.filter(project_id__in=allowed_pids)

    if date_from:
        qs = qs.filter(report_date__gte=date_from)
    if date_to:
        qs = qs.filter(report_date__lte=date_to)

    row_count = qs.count()
    if row_count > ROW_LIMIT_CAP:
        raise RowLimitExceededError()

    dprs = list(qs.order_by("-report_date"))

    wb = openpyxl.Workbook()

    # Sheet 1: DPR Log
    ws1 = wb.active
    ws1.title = "DPR Summary Log"

    ws1.cell(row=1, column=1, value=f"{org.name} - Daily Progress Report Log").font = Font(
        name="Arial", size=14, bold=True
    )
    ws1.cell(
        row=2,
        column=1,
        value=f"Period: {date_from or 'All'} to {date_to or 'All'} | Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
    ).font = Font(name="Arial", size=9, color="475569")

    headers_1 = [
        "Date",
        "Project Code",
        "Project Name",
        "Status",
        "Weather",
        "Work Summary",
        "Issues / Obstacles",
        "Submitted By",
        "Verified By",
    ]
    for c_idx, h in enumerate(headers_1, start=1):
        ws1.cell(row=4, column=c_idx, value=h)

    current_row = 5
    for d in dprs:
        ws1.cell(row=current_row, column=1, value=str(d.report_date))
        ws1.cell(row=current_row, column=2, value=d.project.code)
        ws1.cell(row=current_row, column=3, value=d.project.name)
        ws1.cell(row=current_row, column=4, value=d.status)
        ws1.cell(row=current_row, column=5, value=d.weather or "N/A")
        ws1.cell(row=current_row, column=6, value=d.summary)
        ws1.cell(row=current_row, column=7, value=d.issues or "")
        ws1.cell(row=current_row, column=8, value=get_user_display_name(d.submitted_by))
        ws1.cell(
            row=current_row,
            column=9,
            value=get_user_display_name(d.verified_by) if d.verified_by else "Pending",
        )
        current_row += 1

    apply_excel_styling(ws1, header_row_index=4)

    # Sheet 2: Tasks and Materials
    ws2 = wb.create_sheet(title="Tasks & Materials")
    ws2.cell(row=1, column=1, value="DPR Tasks & Consumed Materials").font = Font(
        name="Arial", size=14, bold=True
    )

    headers_2 = [
        "DPR Date",
        "Project Code",
        "Type",
        "Item Name",
        "Progress % / Quantity",
        "Unit / Notes",
    ]
    for c_idx, h in enumerate(headers_2, start=1):
        ws2.cell(row=4, column=c_idx, value=h)

    curr_row_2 = 5
    for d in dprs:
        for t in d.tasks.all():
            ws2.cell(row=curr_row_2, column=1, value=str(d.report_date))
            ws2.cell(row=curr_row_2, column=2, value=d.project.code)
            ws2.cell(row=curr_row_2, column=3, value="TASK")
            ws2.cell(row=curr_row_2, column=4, value=t.task_name)
            ws2.cell(row=curr_row_2, column=5, value=float(t.percent_complete))
            ws2.cell(row=curr_row_2, column=6, value=t.notes or "% Complete")
            curr_row_2 += 1

        for m in d.materials.all():
            ws2.cell(row=curr_row_2, column=1, value=str(d.report_date))
            ws2.cell(row=curr_row_2, column=2, value=d.project.code)
            ws2.cell(row=curr_row_2, column=3, value="MATERIAL")
            ws2.cell(row=curr_row_2, column=4, value=m.material_name)
            ws2.cell(row=curr_row_2, column=5, value=float(m.quantity))
            ws2.cell(row=curr_row_2, column=6, value=m.unit)
            curr_row_2 += 1

    apply_excel_styling(ws2, header_row_index=4)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, row_count


# ==============================================================================
# 3. Itemized Expenses Excel Export
# ==============================================================================


def generate_expenses_excel(
    org,
    user,
    project_id=None,
    date_from=None,
    date_to=None,
    category=None,
    status=None,
    allowed_pids=None,
) -> tuple[BytesIO, int]:
    """
    Generate itemized expense ledger spreadsheet with totals row.
    """
    qs = Expense.objects.filter(organization=org, deleted_at__isnull=True).select_related(
        "project", "vendor", "created_by", "approved_by"
    )

    if project_id:
        qs = qs.filter(project_id=project_id)
    elif allowed_pids is not None:
        qs = qs.filter(project_id__in=allowed_pids)

    if date_from:
        qs = qs.filter(expense_date__gte=date_from)
    if date_to:
        qs = qs.filter(expense_date__lte=date_to)
    if category:
        qs = qs.filter(category=category)
    if status:
        qs = qs.filter(status=status)

    row_count = qs.count()
    if row_count > ROW_LIMIT_CAP:
        raise RowLimitExceededError()

    expenses = list(qs.order_by("-expense_date"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expense Ledger"

    ws.cell(row=1, column=1, value=f"{org.name} - Itemized Expense Ledger").font = Font(
        name="Arial", size=14, bold=True
    )
    ws.cell(
        row=2,
        column=1,
        value=f"Filters: Dates [{date_from or 'All'} to {date_to or 'All'}], Category [{category or 'All'}], Status [{status or 'All'}] | "
        f"Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
    ).font = Font(name="Arial", size=9, color="475569")

    headers = [
        "Expense Date",
        "Project Code",
        "Project Name",
        "Category",
        "Description",
        "Amount (INR)",
        "Payment Mode",
        "Vendor",
        "Status",
        "Approved By",
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=4, column=c_idx, value=h)

    curr_row = 5
    total_amount = Decimal("0.00")
    for exp in expenses:
        total_amount += exp.amount
        ws.cell(row=curr_row, column=1, value=str(exp.expense_date))
        ws.cell(row=curr_row, column=2, value=exp.project.code)
        ws.cell(row=curr_row, column=3, value=exp.project.name)
        ws.cell(row=curr_row, column=4, value=exp.category)
        ws.cell(row=curr_row, column=5, value=exp.description)
        amt_cell = ws.cell(row=curr_row, column=6, value=float(exp.amount))
        amt_cell.number_format = "#,##0.00"
        amt_cell.alignment = Alignment(horizontal="right")
        ws.cell(row=curr_row, column=7, value=exp.payment_mode)
        ws.cell(row=curr_row, column=8, value=exp.vendor.name if exp.vendor else "N/A")
        ws.cell(row=curr_row, column=9, value=exp.status)
        ws.cell(
            row=curr_row,
            column=10,
            value=get_user_display_name(exp.approved_by) if exp.approved_by else "N/A",
        )
        curr_row += 1

    # Total Row
    ws.cell(row=curr_row, column=1, value="TOTAL").font = Font(name="Arial", size=10, bold=True)
    tot_cell = ws.cell(row=curr_row, column=6, value=float(total_amount))
    tot_cell.font = Font(name="Arial", size=10, bold=True)
    tot_cell.number_format = "#,##0.00"
    tot_cell.alignment = Alignment(horizontal="right")

    apply_excel_styling(ws, header_row_index=4)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, row_count


# ==============================================================================
# 4. Attendance Muster Roll Excel Export
# ==============================================================================


def generate_attendance_excel(
    org, user, project_id=None, date_from=None, date_to=None, allowed_pids=None
) -> tuple[BytesIO, int]:
    """
    Generate worker attendance muster roll matrix and daily log.
    """
    qs = Attendance.objects.filter(organization=org, deleted_at__isnull=True).select_related(
        "employee", "project"
    )

    if project_id:
        qs = qs.filter(project_id=project_id)
    elif allowed_pids is not None:
        qs = qs.filter(project_id__in=allowed_pids)

    if date_from:
        qs = qs.filter(date__gte=date_from)
    if date_to:
        qs = qs.filter(date__lte=date_to)

    row_count = qs.count()
    if row_count > ROW_LIMIT_CAP:
        raise RowLimitExceededError()

    wb = openpyxl.Workbook()

    # Sheet 1: Worker Summary Muster Roll
    ws1 = wb.active
    ws1.title = "Muster Roll Summary"

    ws1.cell(row=1, column=1, value=f"{org.name} - Worker Attendance Muster Roll").font = Font(
        name="Arial", size=14, bold=True
    )
    ws1.cell(
        row=2,
        column=1,
        value=f"Period: {date_from or 'All'} to {date_to or 'All'} | Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
    ).font = Font(name="Arial", size=9, color="475569")

    headers_1 = [
        "Worker Code",
        "Worker Name",
        "Trade",
        "Days Present",
        "Half Days",
        "Days Absent",
        "Total OT Hours",
    ]
    for c_idx, h in enumerate(headers_1, start=1):
        ws1.cell(row=4, column=c_idx, value=h)

    # Worker aggregation
    summary_by_emp = {}
    attendance_records = list(qs.order_by("date"))

    for a in attendance_records:
        emp_id = a.employee_id
        if emp_id not in summary_by_emp:
            summary_by_emp[emp_id] = {
                "code": a.employee.employee_code,
                "name": a.employee.full_name,
                "trade": a.employee.trade,
                "present": 0,
                "half": 0,
                "absent": 0,
                "ot_hours": Decimal("0.00"),
            }

        if a.status == AttendanceStatus.PRESENT:
            summary_by_emp[emp_id]["present"] += 1
        elif a.status == AttendanceStatus.HALF_DAY:
            summary_by_emp[emp_id]["half"] += 1
        elif a.status == AttendanceStatus.ABSENT:
            summary_by_emp[emp_id]["absent"] += 1

        summary_by_emp[emp_id]["ot_hours"] += a.ot_hours

    curr_row = 5
    tot_pres = 0
    tot_half = 0
    tot_abs = 0
    tot_ot = Decimal("0.00")

    for emp_data in summary_by_emp.values():
        tot_pres += emp_data["present"]
        tot_half += emp_data["half"]
        tot_abs += emp_data["absent"]
        tot_ot += emp_data["ot_hours"]

        ws1.cell(row=curr_row, column=1, value=emp_data["code"])
        ws1.cell(row=curr_row, column=2, value=emp_data["name"])
        ws1.cell(row=curr_row, column=3, value=emp_data["trade"])
        ws1.cell(row=curr_row, column=4, value=emp_data["present"])
        ws1.cell(row=curr_row, column=5, value=emp_data["half"])
        ws1.cell(row=curr_row, column=6, value=emp_data["absent"])
        ws1.cell(row=curr_row, column=7, value=float(emp_data["ot_hours"]))
        curr_row += 1

    # Totals Row
    ws1.cell(row=curr_row, column=1, value="TOTALS").font = Font(name="Arial", size=10, bold=True)
    ws1.cell(row=curr_row, column=4, value=tot_pres).font = Font(name="Arial", size=10, bold=True)
    ws1.cell(row=curr_row, column=5, value=tot_half).font = Font(name="Arial", size=10, bold=True)
    ws1.cell(row=curr_row, column=6, value=tot_abs).font = Font(name="Arial", size=10, bold=True)
    ws1.cell(row=curr_row, column=7, value=float(tot_ot)).font = Font(
        name="Arial", size=10, bold=True
    )

    apply_excel_styling(ws1, header_row_index=4)

    # Sheet 2: Daily Records
    ws2 = wb.create_sheet(title="Daily Attendance Records")
    headers_2 = [
        "Date",
        "Worker Code",
        "Worker Name",
        "Project Code",
        "Status",
        "OT Hours",
        "Locked?",
    ]
    for c_idx, h in enumerate(headers_2, start=1):
        ws2.cell(row=4, column=c_idx, value=h)

    curr_row_2 = 5
    for a in attendance_records:
        ws2.cell(row=curr_row_2, column=1, value=str(a.date))
        ws2.cell(row=curr_row_2, column=2, value=a.employee.employee_code)
        ws2.cell(row=curr_row_2, column=3, value=a.employee.full_name)
        ws2.cell(row=curr_row_2, column=4, value=a.project.code)
        ws2.cell(row=curr_row_2, column=5, value=a.status)
        ws2.cell(row=curr_row_2, column=6, value=float(a.ot_hours))
        ws2.cell(row=curr_row_2, column=7, value="YES" if a.is_locked else "NO")
        curr_row_2 += 1

    apply_excel_styling(ws2, header_row_index=4)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, row_count


# ==============================================================================
# 5. Profit & Loss PDF Export
# ==============================================================================


def generate_profit_loss_pdf(
    org, user, date_from=None, date_to=None, project_id=None
) -> tuple[BytesIO, int]:
    """
    Generate professional Profit & Loss statement PDF from ledger statements.
    """
    pnl_data = profit_and_loss(org=org, date_from=date_from, date_to=date_to, project=project_id)

    project_name = "All Projects"
    if project_id:
        try:
            project_name = Project.objects.get(organization=org, id=project_id).name
        except Project.DoesNotExist:
            project_name = f"Project #{project_id}"

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=142,
        bottomMargin=65,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PnlTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
    )
    meta_style = ParagraphStyle(
        "PnlMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569"),
    )
    section_style = ParagraphStyle(
        "PnlSection",
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1E3A8A"),
    )
    cell_style = ParagraphStyle(
        "PnlCell",
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold = ParagraphStyle(
        "PnlCellBold",
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#0F172A"),
    )

    story = [
        Paragraph(f"{org.name}", title_style),
        Paragraph("PROFIT & LOSS STATEMENT", section_style),
        Spacer(1, 4),
        Paragraph(
            f"Period: {date_from or 'Beginning'} to {date_to or 'Present'} | Project: {project_name}<br/>"
            f"Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
            meta_style,
        ),
        Spacer(1, 14),
    ]

    # Section 1: Revenue
    story.append(Paragraph("1. OPERATING REVENUE", section_style))
    story.append(Spacer(1, 4))

    rev_rows = [
        [
            Paragraph("Account Code", cell_bold),
            Paragraph("Account Name", cell_bold),
            Paragraph("Amount (INR)", cell_bold),
        ]
    ]
    for r in pnl_data.get("revenue_accounts", []):
        rev_rows.append(
            [
                Paragraph(r["account_code"], cell_style),
                Paragraph(r["account_name"], cell_style),
                Paragraph(r["amount"], cell_style),
            ]
        )
    rev_rows.append(
        [
            Paragraph("Total Revenue", cell_bold),
            Paragraph("", cell_bold),
            Paragraph(pnl_data.get("total_revenue", "0.00"), cell_bold),
        ]
    )

    t_rev = Table(rev_rows, colWidths=[100, 280, 150])
    t_rev.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0F172A")),
                ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#0F172A")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(t_rev)
    story.append(Spacer(1, 14))

    # Section 2: Expenses
    story.append(Paragraph("2. OPERATING EXPENSES", section_style))
    story.append(Spacer(1, 4))

    exp_rows = [
        [
            Paragraph("Account Code", cell_bold),
            Paragraph("Account Name", cell_bold),
            Paragraph("Amount (INR)", cell_bold),
        ]
    ]
    for e in pnl_data.get("expense_accounts", []):
        exp_rows.append(
            [
                Paragraph(e["account_code"], cell_style),
                Paragraph(e["account_name"], cell_style),
                Paragraph(e["amount"], cell_style),
            ]
        )
    exp_rows.append(
        [
            Paragraph("Total Expenses", cell_bold),
            Paragraph("", cell_bold),
            Paragraph(pnl_data.get("total_expenses", "0.00"), cell_bold),
        ]
    )

    t_exp = Table(exp_rows, colWidths=[100, 280, 150])
    t_exp.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0F172A")),
                ("LINEABOVE", (0, -1), (-1, -1), 1, colors.HexColor("#0F172A")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(t_exp)
    story.append(Spacer(1, 16))

    # Summary Block: Net Profit
    net_val = Decimal(pnl_data.get("net_profit", "0.00"))
    net_color = "#15803D" if net_val >= Decimal("0.00") else "#B91C1C"

    net_style = ParagraphStyle(
        "NetStyle",
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=14,
        textColor=colors.HexColor(net_color),
    )

    summary_rows = [
        [
            Paragraph("NET PROFIT / (LOSS)", net_style),
            Paragraph(f"INR {pnl_data.get('net_profit', '0.00')}", net_style),
        ]
    ]
    t_sum = Table(summary_rows, colWidths=[380, 150])
    t_sum.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 1.5, colors.HexColor(net_color)),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.append(t_sum)
    doc.build(
        story,
        canvasmaker=NumberedCanvas,
        onFirstPage=draw_letterhead_background,
        onLaterPages=draw_letterhead_background,
    )
    buffer.seek(0)
    return buffer, len(pnl_data.get("revenue_accounts", [])) + len(
        pnl_data.get("expense_accounts", [])
    )


# ==============================================================================
# 6. Trial Balance PDF Export
# ==============================================================================


def generate_trial_balance_pdf(org, user, as_of=None, project_id=None) -> tuple[BytesIO, int]:
    """
    Generate double-entry Trial Balance verification statement PDF.
    """
    tb_data = trial_balance(org=org, as_of=as_of, project=project_id)

    project_name = "All Projects"
    if project_id:
        try:
            project_name = Project.objects.get(organization=org, id=project_id).name
        except Project.DoesNotExist:
            project_name = f"Project #{project_id}"

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=142,
        bottomMargin=65,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TbTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
    )
    section_style = ParagraphStyle(
        "TbSection",
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1E3A8A"),
    )
    meta_style = ParagraphStyle(
        "TbMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#475569"),
    )
    cell_style = ParagraphStyle(
        "TbCell",
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold = ParagraphStyle(
        "TbCellBold",
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#0F172A"),
    )

    is_balanced = tb_data.get("is_balanced", False)
    status_label = "BALANCED" if is_balanced else "UNBALANCED"
    status_color = "#15803D" if is_balanced else "#DC2626"

    story = [
        Paragraph(f"{org.name}", title_style),
        Paragraph("TRIAL BALANCE STATEMENT", section_style),
        Spacer(1, 4),
        Paragraph(
            f"As of: {as_of or 'Current Date'} | Project: {project_name}<br/>"
            f"Status: <font color='{status_color}'><b>{status_label}</b></font> | "
            f"Generated by: {get_user_display_name(user)} on {timezone.now().strftime('%Y-%m-%d %H:%M')}",
            meta_style,
        ),
        Spacer(1, 14),
    ]

    tb_rows = [
        [
            Paragraph("Account Code", cell_bold),
            Paragraph("Account Name", cell_bold),
            Paragraph("Type", cell_bold),
            Paragraph("Debit (INR)", cell_bold),
            Paragraph("Credit (INR)", cell_bold),
        ]
    ]

    for a in tb_data.get("accounts", []):
        tb_rows.append(
            [
                Paragraph(a["account_code"], cell_style),
                Paragraph(a["account_name"], cell_style),
                Paragraph(a["account_type"], cell_style),
                Paragraph(a["debit"], cell_style),
                Paragraph(a["credit"], cell_style),
            ]
        )

    # Grand Totals Row
    tb_rows.append(
        [
            Paragraph("TOTALS", cell_bold),
            Paragraph("", cell_bold),
            Paragraph("", cell_bold),
            Paragraph(tb_data.get("total_debit", "0.00"), cell_bold),
            Paragraph(tb_data.get("total_credit", "0.00"), cell_bold),
        ]
    )

    col_widths = [80, 200, 70, 90, 90]
    table = Table(tb_rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0F172A")),
                ("LINEABOVE", (0, -1), (-1, -1), 1.5, colors.HexColor("#0F172A")),
                ("LINEBELOW", (0, -1), (-1, -1), 1.5, colors.HexColor("#0F172A")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    story.append(table)
    doc.build(
        story,
        canvasmaker=NumberedCanvas,
        onFirstPage=draw_letterhead_background,
        onLaterPages=draw_letterhead_background,
    )
    buffer.seek(0)
    return buffer, len(tb_data.get("accounts", []))
