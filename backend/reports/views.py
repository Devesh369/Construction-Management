"""
Reports generation views.
Exports PDF and Excel sheets.
"""

from io import BytesIO

from django.http import HttpResponse
from rest_framework import views
from rest_framework.permissions import IsAuthenticated

from core.permissions import IsTenantUser, roles_required
from core.roles import Role


class ReportPayrollPDFView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.ACCOUNTANT, Role.ORG_ADMIN),
    ]

    def post(self, request):
        buffer = BytesIO()
        buffer.write(b"%PDF-1.4\n%EOF\n")
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="payroll_report.pdf"'
        return response


class ReportDPRExcelView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.PROJECT_MANAGER, Role.ORG_ADMIN),
    ]

    def post(self, request):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DPR Summary"
        ws.append(["Project", "Date", "Summary", "Status"])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="dpr_report.xlsx"'
        return response


class ReportExpensesExcelView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.ACCOUNTANT, Role.ORG_ADMIN, Role.PROJECT_MANAGER),
    ]

    def post(self, request):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Expenses"
        ws.append(["Project", "Category", "Amount", "Status", "Date"])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="expenses_report.xlsx"'
        return response


class ReportProfitLossPDFView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.ACCOUNTANT, Role.ORG_ADMIN),
    ]

    def post(self, request):
        buffer = BytesIO()
        buffer.write(b"%PDF-1.4\n%EOF\n")
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="profit_loss.pdf"'
        return response


class ReportAttendanceExcelView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.PROJECT_MANAGER, Role.ORG_ADMIN, Role.ACCOUNTANT),
    ]

    def post(self, request):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Matrix"
        ws.append(["Worker Code", "Name", "Trade", "Days Present", "OT Hours"])
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="attendance_report.xlsx"'
        return response


class ReportTrialBalancePDFView(views.APIView):
    permission_classes = [
        IsAuthenticated,
        IsTenantUser,
        roles_required(Role.ACCOUNTANT, Role.ORG_ADMIN),
    ]

    def post(self, request):
        buffer = BytesIO()
        buffer.write(b"%PDF-1.4\n%EOF\n")
        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="trial_balance.pdf"'
        return response
